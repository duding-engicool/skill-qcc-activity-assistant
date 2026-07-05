#!/usr/bin/env python3
"""
PPT渲染脚本
将大纲YAML渲染为.pptx文件
使用python-pptx直接生成，朴素工业风，优化排版
"""

import argparse
import json
import sys
import os
from pathlib import Path

try:
    import yaml
    from pptx import Presentation
    from pptx.util import Inches, Pt, Cm, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
except ImportError as e:
    print(json.dumps({
        "status": "error",
        "message": f"依赖缺失: {e}，请运行: pip install python-pptx pyyaml"
    }, ensure_ascii=False))
    sys.exit(1)


# 配色方案（朴素工业风）
COLORS = {
    "primary_blue": RGBColor(0x1F, 0x4E, 0x79),  # 主蓝
    "text_gray": RGBColor(0x59, 0x59, 0x59),      # 正文灰
    "accent_orange": RGBColor(0xED, 0x7D, 0x31),  # 强调橙
    "warning_red": RGBColor(0xC0, 0x00, 0x00),    # 警示红
    "light_gray": RGBColor(0xF2, 0xF2, 0xF2),     # 浅灰底
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "black": RGBColor(0x00, 0x00, 0x00),
    "border_gray": RGBColor(0xBD, 0xBD, 0xBD)    # 边框灰
}

# 布局常量
SLIDE_WIDTH_CM = 33.87  # 16:9 宽度
SLIDE_HEIGHT_CM = 19.05  # 16:9 高度
HEADER_HEIGHT_CM = 1.5
CONTENT_TOP_CM = 2.0
MARGIN_CM = 1.5


def add_header_bar(slide, title_text, prs):
    """添加页眉蓝色条"""
    # 顶部蓝色条
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 
        0, 0, 
        prs.slide_width, 
        Cm(HEADER_HEIGHT_CM)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["primary_blue"]
    shape.line.fill.background()
    
    # 标题文字
    txBox = slide.shapes.add_textbox(
        Cm(MARGIN_CM), 
        Cm(0.3), 
        Cm(SLIDE_WIDTH_CM - 2 * MARGIN_CM), 
        Cm(0.9)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(22)
    p.font.color.rgb = COLORS["white"]
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT


def add_text_block(slide, block, left, top, width, height):
    """添加文字块，优化排版"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    block_type = block.get("type", "paragraph")
    font_size = block.get("font_size", 14)
    
    if block_type == "bullet":
        items = block.get("items", [])
        for i, item in enumerate(items):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(font_size)
            p.font.color.rgb = COLORS["text_gray"]
            p.space_before = Pt(4)
            p.space_after = Pt(4)
            p.level = 0
    
    elif block_type == "paragraph":
        content = block.get("content", "")
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["text_gray"]
        p.line_spacing = 1.3
    
    elif block_type == "highlight":
        # 高亮文本（用于关键数据）
        content = block.get("content", "")
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["accent_orange"]
        p.font.bold = True
    
    elif block_type == "image_analysis":
        # 图片解析文字
        content = block.get("content", "")
        p = tf.paragraphs[0]
        p.text = f"【图析】{content}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLORS["primary_blue"]
        p.font.italic = True
        p.line_spacing = 1.2
    
    return txBox


def add_table(slide, block, left, top, width, height):
    """添加表格，优化样式"""
    headers = block.get("headers", [])
    rows = block.get("rows", [])
    
    if not headers:
        return None
    
    num_rows = len(rows) + 1
    num_cols = len(headers)
    
    # 计算合适的行高
    row_height = min(Cm(1.0), height // num_rows)
    
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table
    
    # 设置列宽均匀分布
    col_width = width // num_cols
    for i in range(num_cols):
        table.columns[i].width = col_width
    
    # 设置表头
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = str(header)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["primary_blue"]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.color.rgb = COLORS["white"]
            paragraph.font.size = Pt(11)
            paragraph.font.bold = True
            paragraph.alignment = PP_ALIGN.CENTER
    
    # 设置数据行
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_data in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(cell_data)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(10)
                paragraph.font.color.rgb = COLORS["text_gray"]
                paragraph.alignment = PP_ALIGN.CENTER
            
            # 隔行变色
            if row_idx % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["light_gray"]
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["white"]
    
    # 设置边框
    for row in table.rows:
        for cell in row.cells:
            for border in cell.border:
                border.color.rgb = COLORS["border_gray"]
                border.width = Pt(0.5)
    
    return table


def add_image_with_analysis(slide, image_info, project_root, left, top, width, height):
    """添加图片及其解析文字"""
    img_path = project_root / image_info.get("path", "")
    analysis = image_info.get("analysis", "")
    
    if not img_path.exists():
        # 图片缺失占位符
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS["light_gray"]
        shape.line.color.rgb = COLORS["warning_red"]
        shape.line.width = Pt(2)
        
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = f"[图片缺失]\n{image_info.get('path', '')}"
        p.font.color.rgb = COLORS["warning_red"]
        p.font.size = Pt(11)
        p.alignment = PP_ALIGN.CENTER
        return False
    
    try:
        # 添加图片
        pic = slide.shapes.add_picture(str(img_path), left, top, width=width)
        
        # 如果有解析文字，在图片下方添加
        if analysis:
            analysis_top = top + pic.height + Cm(0.2)
            analysis_height = Cm(2.0)
            
            txBox = slide.shapes.add_textbox(left, analysis_top, width, analysis_height)
            tf = txBox.text_frame
            tf.word_wrap = True
            
            p = tf.paragraphs[0]
            p.text = f"【图析】{analysis}"
            p.font.size = Pt(10)
            p.font.color.rgb = COLORS["primary_blue"]
            p.font.italic = True
            p.line_spacing = 1.2
        
        return True
    except Exception as e:
        print(f"警告: 无法插入图片 {img_path}: {e}", file=sys.stderr)
        return False


def render_cover_page(slide, page_data, prs):
    """渲染封面页"""
    blocks = page_data.get("blocks", [])
    
    # 封面背景 - 大标题居中
    y_pos = Cm(5)
    for block in blocks:
        font_size = block.get("font_size", 24)
        content = block.get("content", "")
        
        txBox = slide.shapes.add_textbox(
            Cm(3), y_pos, 
            Cm(SLIDE_WIDTH_CM - 6), 
            Cm(2)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["primary_blue"] if font_size >= 20 else COLORS["text_gray"]
        p.font.bold = (font_size >= 20)
        p.alignment = PP_ALIGN.CENTER
        
        y_pos += Cm(2.5)


def render_section_break_page(slide, page_data, prs):
    """渲染章节过渡页"""
    blocks = page_data.get("blocks", [])
    
    # 章节标题居中
    for block in blocks:
        content = block.get("content", "")
        
        txBox = slide.shapes.add_textbox(
            Cm(5), Cm(7), 
            Cm(SLIDE_WIDTH_CM - 10), 
            Cm(3)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(28)
        p.font.color.rgb = COLORS["primary_blue"]
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER


def render_text_image_page(slide, page_data, prs, project_root, image_on_right=True):
    """渲染图文混排页"""
    blocks = page_data.get("blocks", [])
    images = page_data.get("images", [])
    
    # 计算布局
    if image_on_right:
        text_left = Cm(MARGIN_CM)
        text_width = Cm(14)
        img_left = Cm(16)
        img_width = Cm(15)
    else:
        img_left = Cm(MARGIN_CM)
        img_width = Cm(15)
        text_left = Cm(17)
        text_width = Cm(14)
    
    content_top = Cm(CONTENT_TOP_CM)
    content_height = Cm(14)
    
    # 渲染文字块
    y_offset = content_top
    for block in blocks:
        if block.get("type") == "table":
            add_table(slide, block, text_left, y_offset, text_width, Cm(10))
            y_offset += Cm(10.5)
        else:
            block_height = Cm(1.5) * max(1, len(block.get("items", [block.get("content", "")])))
            add_text_block(slide, block, text_left, y_offset, text_width, block_height)
            y_offset += block_height + Cm(0.3)
    
    # 渲染图片
    for img in images:
        add_image_with_analysis(
            slide, img, project_root,
            img_left, content_top, 
            img_width, Cm(12)
        )


def render_table_full_page(slide, page_data, prs):
    """渲染全页表格"""
    blocks = page_data.get("blocks", [])
    
    for block in blocks:
        if block.get("type") == "table":
            add_table(
                slide, block,
                Cm(MARGIN_CM), Cm(CONTENT_TOP_CM),
                Cm(SLIDE_WIDTH_CM - 2 * MARGIN_CM), Cm(14)
            )


def render_image_full_page(slide, page_data, prs, project_root):
    """渲染全屏图片页"""
    images = page_data.get("images", [])
    blocks = page_data.get("blocks", [])
    
    for img in images:
        # 图片居中，稍小一些留出边距
        add_image_with_analysis(
            slide, img, project_root,
            Cm(3), Cm(2.5),
            Cm(SLIDE_WIDTH_CM - 6), Cm(13)
        )


def render_acknowledgement_page(slide, page_data, prs):
    """渲染致谢页"""
    blocks = page_data.get("blocks", [])
    
    for block in blocks:
        content = block.get("content", "")
        
        txBox = slide.shapes.add_textbox(
            Cm(5), Cm(7),
            Cm(SLIDE_WIDTH_CM - 10), Cm(3)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = content
        p.font.size = Pt(24)
        p.font.color.rgb = COLORS["primary_blue"]
        p.alignment = PP_ALIGN.CENTER


def render_text_single_page(slide, page_data, prs):
    """渲染单栏文字页"""
    blocks = page_data.get("blocks", [])
    
    y_offset = Cm(CONTENT_TOP_CM)
    for block in blocks:
        if block.get("type") == "table":
            add_table(
                slide, block,
                Cm(MARGIN_CM), y_offset,
                Cm(SLIDE_WIDTH_CM - 2 * MARGIN_CM), Cm(12)
            )
            y_offset += Cm(12.5)
        else:
            items = block.get("items", [block.get("content", "")])
            block_height = Cm(1.2) * len(items)
            add_text_block(
                slide, block,
                Cm(MARGIN_CM + 1), y_offset,
                Cm(SLIDE_WIDTH_CM - 2 * MARGIN_CM - 2), block_height
            )
            y_offset += block_height + Cm(0.5)


def render_page(slide, page_data, prs, project_root):
    """渲染单页，根据layout类型调度"""
    layout = page_data.get("layout", "text_single")
    title = page_data.get("title", "")
    
    # 添加页眉
    add_header_bar(slide, title, prs)
    
    # 根据layout类型渲染
    if layout == "cover":
        render_cover_page(slide, page_data, prs)
    elif layout == "section_break":
        render_section_break_page(slide, page_data, prs)
    elif layout == "text_left_image_right":
        render_text_image_page(slide, page_data, prs, project_root, image_on_right=True)
    elif layout == "image_left_text_right":
        render_text_image_page(slide, page_data, prs, project_root, image_on_right=False)
    elif layout == "table_full":
        render_table_full_page(slide, page_data, prs)
    elif layout == "image_full":
        render_image_full_page(slide, page_data, prs, project_root)
    elif layout == "acknowledgement":
        render_acknowledgement_page(slide, page_data, prs)
    else:  # text_single
        render_text_single_page(slide, page_data, prs)


def build_pptx(outline_path, master_path, output_path, project_root):
    """构建PPT"""
    # 加载大纲
    with open(outline_path, 'r', encoding='utf-8') as f:
        outline = yaml.safe_load(f)
    
    pages = outline.get("pages", [])
    
    # 检查页数
    if len(pages) > 100:
        print(json.dumps({
            "status": "error",
            "message": f"总页数 {len(pages)} 超过100页上限，请精简大纲"
        }, ensure_ascii=False))
        sys.exit(1)
    
    # 创建演示文稿
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # 渲染每页
    render_log = []
    for page_data in pages:
        slide = prs.slides.add_slide(blank_layout)
        render_page(slide, page_data, prs, project_root)
        
        # 记录渲染日志
        images = page_data.get("images", [])
        has_analysis = any(img.get("analysis") for img in images)
        
        log_entry = {
            "page_no": page_data.get("page_no"),
            "title": page_data.get("title"),
            "layout": page_data.get("layout"),
            "image_count": len(images),
            "has_analysis": has_analysis,
            "status": "✅"
        }
        render_log.append(log_entry)
    
    # 保存
    prs.save(output_path)
    
    return render_log


def main():
    parser = argparse.ArgumentParser(description='将大纲渲染为PPT')
    parser.add_argument('--outline', required=True, help='大纲YAML路径')
    parser.add_argument('--master', default=None, help='母版路径(可选)')
    parser.add_argument('--output', default='发表稿.pptx', help='输出PPT路径')
    parser.add_argument('--project-root', default='.', help='项目根目录(用于解析图片路径)')
    
    args = parser.parse_args()
    
    outline_path = Path(args.outline)
    if not outline_path.exists():
        print(json.dumps({
            "status": "error",
            "message": f"大纲文件不存在: {args.outline}"
        }, ensure_ascii=False))
        sys.exit(1)
    
    project_root = Path(args.project_root)
    output_path = Path(args.output)
    
    # 构建PPT
    render_log = build_pptx(outline_path, args.master, output_path, project_root)
    
    # 输出结果
    total_images = sum(entry["image_count"] for entry in render_log)
    analyzed_images = sum(1 for entry in render_log if entry["has_analysis"])
    
    result = {
        "status": "success",
        "message": f"PPT已生成: {output_path}",
        "data": {
            "total_pages": len(render_log),
            "total_images": total_images,
            "analyzed_images": analyzed_images,
            "output_file": str(output_path)
        }
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
